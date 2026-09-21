# -*- coding: utf-8 -*-
"""
Interfaz gráfica del sistema de extracción de comprobantes de pago.
==================================================================
Reutiliza la lógica de procesar_comprobantes.py (Gemini Flash, validaciones
y escritura del archivo Excel .xlsx). La ruta del Excel se toma de la línea
RUTA_EXCEL de config.txt (puede ser una carpeta de Google Drive).

Todo se maneja desde la interfaz: la clave de API y la ruta del Excel se
configuran en una pantalla de bienvenida (o con el botón ⚙️ Configuración),
las imágenes se agregan una a una, por carpetas o con la cámara web, y el
historial de comprobantes se lee directamente del Excel.

Construida con customtkinter. Se abre con doble clic en ABRIR.bat
(Windows) o ABRIR.command (macOS). También funciona empaquetada con
PyInstaller (--onefile --windowed): las rutas se resuelven respecto a la
carpeta del ejecutable.
"""

import logging
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import webbrowser
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk


# ---------------------------------------------------------------------------
# Carpeta del programa (funciona como .py y empaquetado con PyInstaller)
# ---------------------------------------------------------------------------
def base_dir():
    """Carpeta donde viven config.txt, entrada/, procesados/, etc.
    Con PyInstaller (--onefile) es la carpeta del ejecutable; en caso
    contrario, la carpeta de este archivo .py."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


BASE_DIR = base_dir()
os.chdir(BASE_DIR)
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)


# ---------------------------------------------------------------------------
# Registro de errores técnicos (el usuario solo ve mensajes en español)
# ---------------------------------------------------------------------------
ARCHIVO_LOG = os.path.join(BASE_DIR, "errores.log")


def _configurar_log():
    registro = logging.getLogger("comprobantes")
    registro.setLevel(logging.INFO)
    registro.propagate = False
    try:
        manejador = logging.FileHandler(ARCHIVO_LOG, encoding="utf-8", delay=True)
        manejador.setFormatter(
            logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
        registro.addHandler(manejador)
    except OSError:
        registro.addHandler(logging.NullHandler())
    return registro


LOG = _configurar_log()


def registrar_error(contexto):
    """Guarda en errores.log el error que se está manejando (con traceback)."""
    try:
        LOG.exception(contexto)
    except Exception:
        pass


def _mostrar_error_inesperado():
    try:
        messagebox.showerror(
            "Error inesperado",
            "Ocurrió un error inesperado. El programa puede seguir funcionando.\n\n"
            "Si el problema se repite, envía el archivo errores.log "
            "(está en la carpeta del programa) para revisarlo.",
        )
    except Exception:
        pass


_app_actual = None  # ventana principal, para avisar errores desde otros hilos


def _excepthook(tipo, valor, tb):
    if issubclass(tipo, KeyboardInterrupt):
        return
    try:
        LOG.error("Error no controlado", exc_info=(tipo, valor, tb))
    except Exception:
        pass
    _mostrar_error_inesperado()


def _excepthook_hilo(args):
    try:
        LOG.error("Error no controlado en un hilo",
                  exc_info=(args.exc_type, args.exc_value, args.exc_traceback))
    except Exception:
        pass
    # Los cuadros de diálogo solo pueden mostrarse desde el hilo principal
    if _app_actual is not None:
        try:
            _app_actual.after(0, _mostrar_error_inesperado)
        except Exception:
            pass


sys.excepthook = _excepthook
threading.excepthook = _excepthook_hilo


# ---------------------------------------------------------------------------
# Dependencias
# ---------------------------------------------------------------------------
try:
    import customtkinter as ctk
except ImportError:
    ctk = None

# Lógica compartida con el procesador de consola
try:
    from procesar_comprobantes import (
        ARCHIVO_CONFIG,
        BASE,
        CARPETA_ENTRADA,
        CARPETA_PROCESADOS,
        CARPETA_REVISION,
        COLUMNAS,
        EXTENSIONES,
        MODELO,
        PAUSA_SEGUNDOS,
        RUTA_EXCEL_POR_DEFECTO,
        ExcelBloqueadoError,
        agregar_y_guardar,
        cargar_comprobantes_existentes,
        construir_fila,
        crear_carpetas,
        extraer_datos,
        leer_ruta_excel,
        listar_imagenes,
        mapear_columnas,
        mover_imagen,
        preparar_imagen,
        validar,
    )
except (ImportError, SystemExit, RuntimeError, EOFError):
    # procesar_comprobantes.py termina el programa si faltan dependencias
    # (google-genai, Pillow u openpyxl). Sin consola solo podemos avisar así.
    registrar_error("No se pudo importar procesar_comprobantes")
    try:
        raiz = tk.Tk()
        raiz.withdraw()
        messagebox.showerror(
            "Faltan dependencias",
            "Faltan componentes del sistema para poder abrir el programa.\n\n"
            "Solución: haz doble clic en INSTALAR.bat (o INSTALAR.command en Mac), "
            "espera a que termine y vuelve a abrir el programa.",
        )
    except Exception:
        pass
    sys.exit(1)

if ctk is None:
    try:
        raiz = tk.Tk()
        raiz.withdraw()
        messagebox.showerror(
            "Faltan dependencias",
            "No está instalada la librería customtkinter.\n\n"
            "Solución: haz doble clic en INSTALAR.bat (o INSTALAR.command en Mac) "
            "y espera a que termine.",
        )
    except Exception:
        pass
    sys.exit(1)

try:
    from google import genai
except ImportError:
    genai = None

try:
    from PIL import Image, ImageDraw, ImageFont, ImageTk
except ImportError:
    Image = None
    ImageTk = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

# ---------------------------------------------------------------------------
# Apariencia
# ---------------------------------------------------------------------------
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

COLOR_FONDO = "#0f1117"
COLOR_PANEL = "#171a23"
COLOR_TARJETA = "#212532"
COLOR_TARJETA_HOVER = "#2a2f3e"
COLOR_BORDE = "#2c3140"
COLOR_TEXTO = "#f5f6fa"
COLOR_TEXTO_SUAVE = "#8b90a0"
COLOR_ACENTO = "#e94560"
COLOR_ACENTO_HOVER = "#c23650"
COLOR_DESHABILITADO = "#3a3f4d"
COLOR_OK = "#22c55e"
COLOR_REVISAR = "#facc15"
COLOR_ERROR = "#ef4444"
COLOR_SELECCION = "#253562"
COLOR_SELECCION_BORDE = "#3b82f6"
COLOR_ENLACE = "#60a5fa"

INTERVALO_REFRESCO_MS = 3000   # refresco del estado de la API key y de la lista
LARGO_MAXIMO_NOMBRE = 46       # caracteres visibles del nombre en la lista
LIMITE_HISTORIAL = 20          # filas del Excel que se muestran en el historial
ESPERA_CUOTA_SEGUNDOS = 60     # espera antes de reintentar tras un límite de Gemini
URL_API_KEY = "https://aistudio.google.com/apikey"
NOMBRE_EXCEL_POR_DEFECTO = "comprobantes.xlsx"

MENSAJE_EXCEL_NO_LEIDO = ("No se pudo leer el archivo de Excel. "
                          "Verifica que no esté abierto en otro programa.")


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------
def api_key_valida(valor):
    """True si el texto parece una clave configurada (no vacía ni plantilla)."""
    valor = (valor or "").strip()
    return bool(valor) and "PEGA_AQUI" not in valor


def leer_config_gui():
    """Lee config.txt y devuelve un diccionario {CLAVE: valor} sin imprimir
    en consola. Si el archivo no existe o no se puede leer, devuelve {}."""
    valores = {}
    if not ARCHIVO_CONFIG.exists():
        return valores
    try:
        contenido = ARCHIVO_CONFIG.read_text(encoding="utf-8-sig")
    except OSError:
        registrar_error("No se pudo leer config.txt")
        return valores
    for linea in contenido.splitlines():
        linea = linea.strip()
        if linea.startswith("#") or "=" not in linea:
            continue
        clave, valor = linea.split("=", 1)
        valores[clave.strip()] = valor.strip()
    return valores


def leer_api_key_gui():
    """Lee la API key desde config.txt sin imprimir en consola.
    Devuelve la clave o None si falta o no está configurada."""
    valor = leer_config_gui().get("API_KEY", "")
    return valor if api_key_valida(valor) else None


def leer_ruta_excel_config():
    """Devuelve el texto de RUTA_EXCEL tal como está en config.txt ('' si falta)."""
    return leer_config_gui().get("RUTA_EXCEL", "").strip('"').strip("'").strip()


# Plantilla usada solo si no existe config.txt ni config.ejemplo.txt
PLANTILLA_CONFIG = (
    "# Clave de API de Google Gemini (gratis en https://aistudio.google.com/apikey)\n"
    "API_KEY=PEGA_AQUI_TU_CLAVE\n"
    "\n"
    "# Cambia esta ruta para guardar el Excel en Google Drive. Ejemplo: "
    "C:\\Users\\TuUsuario\\Google Drive\\Mi unidad\\Comprobantes\\comprobantes.xlsx\n"
    "RUTA_EXCEL=salida/comprobantes.xlsx\n"
)


def guardar_config(api_key, ruta_excel):
    """Escribe API_KEY y RUTA_EXCEL en config.txt.

    Si config.txt existe, solo se reemplazan esas dos líneas y se conserva
    todo lo demás (comentarios incluidos). Si no existe, se crea a partir de
    config.ejemplo.txt (o de una plantilla equivalente). Lanza OSError si no
    se puede escribir.
    """
    api_key = (api_key or "").strip()
    ruta_excel = (ruta_excel or "").strip() or RUTA_EXCEL_POR_DEFECTO
    pendientes = {"API_KEY": api_key, "RUTA_EXCEL": ruta_excel}

    if ARCHIVO_CONFIG.exists():
        lineas = ARCHIVO_CONFIG.read_text(encoding="utf-8-sig").splitlines()
    else:
        ejemplo = BASE / "config.ejemplo.txt"
        try:
            lineas = ejemplo.read_text(encoding="utf-8-sig").splitlines()
        except OSError:
            lineas = PLANTILLA_CONFIG.splitlines()

    salida = []
    for linea in lineas:
        limpia = linea.strip()
        if limpia and not limpia.startswith("#") and "=" in limpia:
            clave = limpia.split("=", 1)[0].strip()
            if clave in pendientes:
                salida.append(f"{clave}={pendientes.pop(clave)}")
                continue
        salida.append(linea)
    for clave, valor in pendientes.items():
        if salida and salida[-1].strip():
            salida.append("")
        salida.append(f"{clave}={valor}")

    # Se escribe en un archivo temporal y luego se reemplaza, para no dejar
    # config.txt a medias si algo falla.
    temporal = ARCHIVO_CONFIG.with_name(ARCHIVO_CONFIG.name + ".tmp")
    temporal.write_text("\n".join(salida) + "\n", encoding="utf-8")
    os.replace(temporal, ARCHIVO_CONFIG)


def formatear_moneda(valor):
    """Formatea un entero como moneda colombiana: $ 150.000"""
    if valor in (None, ""):
        return "—"
    try:
        return "$ " + f"{int(valor):,}".replace(",", ".")
    except (ValueError, TypeError):
        return str(valor)


def tamano_legible(num_bytes):
    """Convierte un tamaño en bytes a texto: 850 B, 245 KB, 1,4 MB."""
    if num_bytes < 1024:
        return f"{num_bytes} B"
    if num_bytes < 1024 * 1024:
        return f"{num_bytes / 1024:.0f} KB"
    return f"{num_bytes / (1024 * 1024):.1f} MB".replace(".", ",")


def acortar(texto, largo=LARGO_MAXIMO_NOMBRE):
    return texto if len(texto) <= largo else texto[: largo - 1] + "…"


def acortar_centro(texto, largo=80):
    """Acorta un texto largo (p. ej. una ruta) dejando el inicio y el final."""
    if len(texto) <= largo:
        return texto
    mitad = (largo - 1) // 2
    return texto[:mitad] + "…" + texto[-mitad:]


# Fuentes de emoji a color por sistema (ruta, tamaño en que se dibujan)
FUENTES_EMOJI = [
    (os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", "seguiemj.ttf"), 64),
    ("/System/Library/Fonts/Apple Color Emoji.ttc", 160),
    ("/usr/share/fonts/truetype/noto/NotoColorEmoji.ttf", 109),
    ("/usr/share/fonts/noto/NotoColorEmoji.ttf", 109),
]
_iconos_emoji = {}


def icono_emoji(emoji, lado=18):
    """Dibuja un emoji a color como CTkImage (Tk los muestra pequeños y sin color).
    Devuelve None si no hay Pillow o fuente de emoji; entonces se usa texto."""
    if emoji in _iconos_emoji:
        return _iconos_emoji[emoji]
    icono = None
    if Image is not None:
        for ruta_fuente, tamano in FUENTES_EMOJI:
            if not os.path.exists(ruta_fuente):
                continue
            try:
                fuente = ImageFont.truetype(ruta_fuente, tamano)
                lienzo = Image.new("RGBA", (tamano * 2, tamano * 2), (0, 0, 0, 0))
                ImageDraw.Draw(lienzo).text((tamano // 2, tamano // 2), emoji,
                                            font=fuente, embedded_color=True)
                caja = lienzo.getbbox()
                if not caja:
                    continue
                recorte = lienzo.crop(caja)
                medida = max(recorte.size)
                cuadro = Image.new("RGBA", (medida, medida), (0, 0, 0, 0))
                cuadro.paste(recorte, ((medida - recorte.width) // 2,
                                       (medida - recorte.height) // 2))
                cuadro = cuadro.resize((lado * 2, lado * 2), Image.LANCZOS)
                icono = ctk.CTkImage(light_image=cuadro, dark_image=cuadro,
                                     size=(lado, lado))
                break
            except Exception:
                continue
    _iconos_emoji[emoji] = icono
    return icono


def enfocar_si_existe(widget, forzar=False):
    """Da el foco a un widget desde un after(); no falla si ya fue destruido."""
    try:
        if widget.winfo_exists():
            widget.focus_force() if forzar else widget.focus_set()
    except Exception:
        pass


def centrar_sobre(ventana, padre, ancho, alto):
    """Coloca una ventana secundaria de 'ancho x alto' centrada sobre 'padre'
    (sin salirse de la pantalla). Si algo falla, solo fija el tamaño."""
    try:
        padre.update_idletasks()
        try:
            escala = ctk.ScalingTracker.get_widget_scaling(padre)
        except Exception:
            escala = 1.0
        ancho_px, alto_px = int(ancho * escala), int(alto * escala)
        x = padre.winfo_rootx() + (padre.winfo_width() - ancho_px) // 2
        y = padre.winfo_rooty() + (padre.winfo_height() - alto_px) // 2
        x = max(0, min(x, padre.winfo_screenwidth() - ancho_px))
        y = max(0, min(y, padre.winfo_screenheight() - alto_px))
        ventana.geometry(f"{ancho}x{alto}+{x}+{y}")
    except Exception:
        ventana.geometry(f"{ancho}x{alto}")


def clasificar_error_gemini(error):
    """Reconoce los fallos más comunes al hablar con Gemini y devuelve:
    'api_key' (clave rechazada), 'cuota' (límite de solicitudes), 'conexion'
    (sin internet o servidor inalcanzable), 'servicio' (Gemini saturado) o
    None si es otro tipo de error."""
    texto = str(error).lower()
    nombre = type(error).__name__.lower()
    # Un PermissionError del sistema (OSError) no es un problema de API key
    if not isinstance(error, OSError) and (
            "api key" in texto or "api_key" in texto
            or "authenticate" in texto or "permission" in texto):
        return "api_key"
    if ("429" in texto or "quota" in texto or "resource" in texto
            or "rate limit" in texto or "too many requests" in texto):
        return "cuota"
    pistas_nombre = ("connect", "timeout", "network", "transport", "remoteprotocol",
                     "readerror", "writeerror", "proxy", "dns", "gaierror")
    pistas_texto = ("getaddrinfo", "name resolution", "nodename", "unreachable",
                    "connection", "timed out", "timeout", "network", "ssl",
                    "temporary failure", "no route to host", "socket")
    if any(p in nombre for p in pistas_nombre) or any(p in texto for p in pistas_texto):
        return "conexion"
    if ("503" in texto or "unavailable" in texto or "overloaded" in texto
            or "deadline" in texto):
        return "servicio"
    return None


def abrir_en_sistema(ruta):
    """Abre un archivo o carpeta con la aplicación predeterminada del sistema."""
    if sys.platform.startswith("win"):
        os.startfile(str(ruta))
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(ruta)])
    else:
        subprocess.Popen(["xdg-open", str(ruta)])


def buscar_imagenes_en_carpeta(carpeta):
    """Devuelve las rutas de todas las imágenes soportadas dentro de la carpeta
    (incluidas subcarpetas), sin contar las que ya están en entrada/."""
    rutas = []
    entrada = os.path.normcase(os.path.abspath(str(CARPETA_ENTRADA)))
    for raiz, _carpetas, archivos in os.walk(carpeta):
        if os.path.normcase(os.path.abspath(raiz)) == entrada:
            continue
        for nombre in sorted(archivos):
            if os.path.splitext(nombre)[1].lower() in EXTENSIONES:
                rutas.append(os.path.join(raiz, nombre))
    return rutas


def nombre_libre(carpeta, nombre):
    """Devuelve una ruta en 'carpeta' que no exista aún, agregando _1, _2..."""
    destino = Path(carpeta) / nombre
    if not destino.exists():
        return destino
    base, extension = os.path.splitext(nombre)
    contador = 1
    while True:
        destino = Path(carpeta) / f"{base}_{contador}{extension}"
        if not destino.exists():
            return destino
        contador += 1


# ---------------------------------------------------------------------------
# Historial: lectura del Excel de resultados
# ---------------------------------------------------------------------------
def _texto_fecha(valor):
    """Convierte el valor de la columna fecha_procesado a texto AAAA-MM-DD HH:MM."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M")
    return str(valor).strip()


def _fecha_orden(texto):
    """Fecha para ordenar el historial; si no se puede interpretar, la mínima."""
    for formato in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto[: len("2000-01-01 00:00:00")].strip(), formato)
        except ValueError:
            continue
    return datetime.min


def leer_historial_excel(ruta, limite=LIMITE_HISTORIAL):
    """Lee el Excel de resultados y devuelve (filas, procesados_hoy, total).

    filas: lista de diccionarios (columnas de COLUMNAS) con los últimos
    'limite' comprobantes, del más reciente al más antiguo, según la columna
    fecha_procesado. Si el Excel no existe devuelve ([], 0, 0).
    """
    ruta = Path(ruta)
    if not ruta.exists() or load_workbook is None:
        return [], 0, 0

    registros = []
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = libro.active
        iterador = hoja.iter_rows(values_only=True)
        encabezados = next(iterador, None) or ()
        # Posición de cada columna según el encabezado (acepta los nombres
        # viejos, p. ej. numero_cuenta). Las que el Excel no tenga quedan vacías.
        indices = mapear_columnas(encabezados)
        for posicion, fila in enumerate(iterador):
            if not fila or all(valor is None for valor in fila):
                continue
            registro = {}
            for columna in COLUMNAS:
                indice = indices.get(columna)
                registro[columna] = (fila[indice] if indice is not None
                                     and indice < len(fila) else None)
            registros.append((posicion, registro))
    finally:
        libro.close()

    hoy = datetime.now().strftime("%Y-%m-%d")
    total = len(registros)
    procesados_hoy = sum(
        1 for _posicion, registro in registros
        if _texto_fecha(registro.get("fecha_procesado")).startswith(hoy)
    )
    registros.sort(
        key=lambda item: (_fecha_orden(_texto_fecha(item[1].get("fecha_procesado"))),
                          item[0]),
        reverse=True,
    )
    return [registro for _posicion, registro in registros[:limite]], procesados_hoy, total


def localizar_imagen_procesada(nombre):
    """Busca la imagen original de un comprobante (columna archivo_origen) en
    procesados/ y, si no está, en revision_manual/. Tiene en cuenta que
    mover_imagen agrega una marca de tiempo si el nombre ya existía.
    Devuelve la ruta (Path) o None."""
    nombre = os.path.basename(str(nombre).strip())
    if not nombre:
        return None
    base, extension = os.path.splitext(nombre)
    patron = re.compile(re.escape(base) + r"_\d{14}" + re.escape(extension) + "$",
                        re.IGNORECASE)
    for carpeta in (CARPETA_PROCESADOS, CARPETA_REVISION):
        candidata = carpeta / nombre
        if candidata.exists():
            return candidata
        try:
            renombradas = sorted(
                (ruta for ruta in carpeta.glob(f"{base}_*{extension}")
                 if patron.match(ruta.name)),
                key=lambda ruta: ruta.name,
            )
        except OSError:
            renombradas = []
        if renombradas:
            return renombradas[-1]
    return None


# ---------------------------------------------------------------------------
# Cámara web (OpenCV es opcional: si no está, el botón queda deshabilitado)
# ---------------------------------------------------------------------------
def _abrir_captura(cv2, indice=0):
    """Abre la cámara. En Windows DirectShow arranca mucho más rápido."""
    try:
        cv2.utils.logging.setLogLevel(cv2.utils.logging.LOG_LEVEL_SILENT)
    except Exception:
        pass
    if sys.platform.startswith("win"):
        return cv2.VideoCapture(indice, cv2.CAP_DSHOW)
    return cv2.VideoCapture(indice)


def detectar_camara():
    """Comprueba si hay cámara. Devuelve (True, '') o (False, motivo).
    Es lenta (puede tardar segundos): llamar desde un hilo."""
    try:
        import cv2
    except ImportError:
        return False, ("Para usar la cámara instala opencv-python\n"
                       "(pip install opencv-python) y vuelve a abrir el programa.")
    except Exception:
        registrar_error("No se pudo cargar OpenCV")
        return False, "No se pudo cargar OpenCV (opencv-python)."
    try:
        captura = _abrir_captura(cv2)
        try:
            disponible = bool(captura.isOpened())
        finally:
            captura.release()
    except Exception:
        registrar_error("Error al detectar la cámara")
        disponible = False
    if disponible:
        return True, ""
    return False, "No se detectó ninguna cámara conectada a este equipo."


class Tooltip:
    """Texto emergente sencillo al pasar el ratón por un widget."""

    def __init__(self, widget, texto):
        self.widget = widget
        self.texto = texto
        self.ventana = None
        self.tarea = None
        widget.bind("<Enter>", self._programar, add="+")
        widget.bind("<Leave>", self._ocultar, add="+")
        widget.bind("<ButtonPress>", self._ocultar, add="+")

    def configurar(self, texto):
        self.texto = texto

    def _programar(self, _evento=None):
        self._cancelar()
        if self.texto:
            self.tarea = self.widget.after(500, self._mostrar)

    def _cancelar(self):
        if self.tarea is not None:
            try:
                self.widget.after_cancel(self.tarea)
            except Exception:
                pass
            self.tarea = None

    def _mostrar(self):
        self.tarea = None
        if self.ventana is not None or not self.texto:
            return
        try:
            x = self.widget.winfo_pointerx() + 14
            y = self.widget.winfo_pointery() + 18
            self.ventana = tk.Toplevel(self.widget)
            self.ventana.wm_overrideredirect(True)
            self.ventana.wm_geometry(f"+{x}+{y}")
            self.ventana.attributes("-topmost", True)
            tk.Label(
                self.ventana, text=self.texto, justify="left",
                bg=COLOR_TARJETA, fg=COLOR_TEXTO, relief="solid", borderwidth=1,
                padx=8, pady=5, wraplength=320,
            ).pack()
        except Exception:
            self.ventana = None

    def _ocultar(self, _evento=None):
        self._cancelar()
        if self.ventana is not None:
            try:
                self.ventana.destroy()
            except Exception:
                pass
            self.ventana = None


class VentanaCamara(ctk.CTkToplevel):
    """Vista previa de la cámara web con botón para capturar.
    La foto se guarda en entrada/ como camara_AAAAMMDD_HHMMSS.jpg."""

    ANCHO_MAXIMO = 640
    ALTO_MAXIMO = 480

    def __init__(self, padre, cv2, al_capturar):
        super().__init__(padre, fg_color=COLOR_FONDO)
        self.cv2 = cv2
        self.al_capturar = al_capturar
        self.detener = threading.Event()
        self.candado = threading.Lock()
        self.ultimo_cuadro = None
        self.foto_tk = None
        self.cerrada = False

        self.title("Cámara web")
        centrar_sobre(self, padre, 700, 620)
        self.minsize(520, 440)
        self.configure(fg_color=COLOR_FONDO)
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self.etiqueta_video = tk.Label(
            self, text="Iniciando la cámara...", bg=COLOR_FONDO, fg=COLOR_TEXTO_SUAVE,
            font=("", 12),
        )
        self.etiqueta_video.grid(row=0, column=0, sticky="nsew", padx=16, pady=(16, 8))

        self.etiqueta_estado = ctk.CTkLabel(
            self, text="Cuando la imagen se vea bien, pulsa Capturar (o la barra espaciadora).",
            text_color=COLOR_TEXTO_SUAVE, font=ctk.CTkFont(size=12),
        )
        self.etiqueta_estado.grid(row=1, column=0, sticky="ew", padx=16)

        botones = ctk.CTkFrame(self, fg_color="transparent")
        botones.grid(row=2, column=0, pady=(8, 16))
        self.boton_capturar = ctk.CTkButton(
            botones, text="📷  Capturar", command=self.capturar, state="disabled",
            font=ctk.CTkFont(size=14, weight="bold"), width=180, height=44,
            corner_radius=22, fg_color=COLOR_ACENTO, hover_color=COLOR_ACENTO_HOVER,
        )
        self.boton_capturar.grid(row=0, column=0, padx=8)
        ctk.CTkButton(
            botones, text="Cerrar", command=self.cerrar, width=120, height=44,
            corner_radius=22, fg_color=COLOR_TARJETA, hover_color=COLOR_TARJETA_HOVER,
            border_width=1, border_color=COLOR_BORDE,
        ).grid(row=0, column=1, padx=8)

        self.bind("<space>", lambda _evento: self.capturar())
        self.bind("<Return>", lambda _evento: self.capturar())
        self.bind("<Escape>", lambda _evento: self.cerrar())
        self.protocol("WM_DELETE_WINDOW", self.cerrar)
        self.after(150, lambda: enfocar_si_existe(self, forzar=True))

        threading.Thread(target=self._leer_en_hilo, daemon=True).start()

    def _en_ui(self, funcion, *args):
        if self.cerrada:
            return
        try:
            self.after(0, funcion, *args)
        except Exception:
            pass

    def _leer_en_hilo(self):
        """Lee cuadros de la cámara en un hilo para no bloquear la interfaz."""
        captura = None
        try:
            captura = _abrir_captura(self.cv2)
            if not captura.isOpened():
                self._en_ui(self._fallo,
                            "No se pudo abrir la cámara. Verifica que no la esté "
                            "usando otro programa.")
                return
            self._en_ui(self._camara_lista)
            fallos = 0
            while not self.detener.is_set():
                ok, cuadro = captura.read()
                if not ok or cuadro is None:
                    fallos += 1
                    if fallos > 100:
                        self._en_ui(self._fallo,
                                    "Se perdió la conexión con la cámara.")
                        return
                    time.sleep(0.05)
                    continue
                fallos = 0
                with self.candado:
                    self.ultimo_cuadro = cuadro
        except Exception:
            registrar_error("Error al leer la cámara")
            self._en_ui(self._fallo, "Ocurrió un problema con la cámara.")
        finally:
            if captura is not None:
                try:
                    captura.release()
                except Exception:
                    pass

    def _camara_lista(self):
        self.boton_capturar.configure(state="normal")
        self._refrescar_cuadro()

    def _fallo(self, mensaje):
        if self.cerrada:
            return
        messagebox.showerror("Cámara web", mensaje, parent=self)
        self.cerrar()

    def _refrescar_cuadro(self):
        if self.cerrada or self.detener.is_set():
            return
        with self.candado:
            cuadro = self.ultimo_cuadro
        if cuadro is not None and Image is not None and ImageTk is not None:
            try:
                rgb = self.cv2.cvtColor(cuadro, self.cv2.COLOR_BGR2RGB)
                imagen = Image.fromarray(rgb)
                disponible_ancho = max(160, self.etiqueta_video.winfo_width() - 8)
                disponible_alto = max(120, self.etiqueta_video.winfo_height() - 8)
                escala = min(disponible_ancho / imagen.width,
                             disponible_alto / imagen.height, 1.0)
                if escala < 1.0:
                    imagen = imagen.resize(
                        (max(1, int(imagen.width * escala)),
                         max(1, int(imagen.height * escala))))
                self.foto_tk = ImageTk.PhotoImage(imagen)
                self.etiqueta_video.configure(image=self.foto_tk, text="")
            except Exception:
                registrar_error("No se pudo mostrar el cuadro de la cámara")
        self.after(40, self._refrescar_cuadro)

    def capturar(self):
        if self.cerrada or self.boton_capturar.cget("state") == "disabled":
            return
        with self.candado:
            cuadro = None if self.ultimo_cuadro is None else self.ultimo_cuadro.copy()
        if cuadro is None:
            self.etiqueta_estado.configure(text="Aún no hay imagen de la cámara. Espera un momento.",
                                           text_color=COLOR_REVISAR)
            return
        try:
            crear_carpetas()
            nombre = "camara_" + datetime.now().strftime("%Y%m%d_%H%M%S") + ".jpg"
            destino = nombre_libre(CARPETA_ENTRADA, nombre)
            rgb = self.cv2.cvtColor(cuadro, self.cv2.COLOR_BGR2RGB)
            if Image is not None:
                Image.fromarray(rgb).save(str(destino), "JPEG", quality=92)
            else:
                ok, datos = self.cv2.imencode(".jpg", cuadro,
                                              [self.cv2.IMWRITE_JPEG_QUALITY, 92])
                if not ok:
                    raise ValueError("No se pudo codificar la imagen")
                destino.write_bytes(datos.tobytes())
        except Exception:
            registrar_error("No se pudo guardar la foto de la cámara")
            messagebox.showerror(
                "Cámara web",
                "No se pudo guardar la foto en la carpeta entrada.\n"
                "Verifica que la carpeta exista y que tengas permiso para escribir en ella.",
                parent=self,
            )
            return
        self.etiqueta_estado.configure(
            text=f"Foto guardada: {destino.name}. Puedes tomar otra o cerrar.",
            text_color=COLOR_OK)
        try:
            self.al_capturar(destino)
        except Exception:
            registrar_error("Error al registrar la foto capturada")

    def cerrar(self):
        if self.cerrada:
            return
        self.cerrada = True
        self.detener.set()
        try:
            self.destroy()
        except Exception:
            pass


# ---------------------------------------------------------------------------
# Formulario de configuración (bienvenida y diálogo ⚙️)
# ---------------------------------------------------------------------------
class FormularioConfiguracion(ctk.CTkFrame):
    """Campos para la clave de API y la ruta del Excel, con botón Guardar.
    'al_guardar' se llama (sin argumentos) cuando config.txt quedó escrito."""

    def __init__(self, padre, al_guardar, texto_boton="Guardar", **opciones):
        opciones.setdefault("fg_color", "transparent")
        super().__init__(padre, **opciones)
        self.al_guardar = al_guardar
        self.grid_columnconfigure(0, weight=1)

        fuente_etiqueta = ctk.CTkFont(size=13, weight="bold")
        fuente_ayuda = ctk.CTkFont(size=11)
        config = leer_config_gui()
        clave_actual = config.get("API_KEY", "")
        if not api_key_valida(clave_actual):
            clave_actual = ""
        ruta_actual = leer_ruta_excel_config()
        if ruta_actual == RUTA_EXCEL_POR_DEFECTO:
            ruta_actual = ""

        # --- Clave de API -------------------------------------------------
        ctk.CTkLabel(self, text="Clave de API de Google Gemini", font=fuente_etiqueta,
                     text_color=COLOR_TEXTO, anchor="w")\
            .grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 4))
        self.entrada_clave = ctk.CTkEntry(
            self, placeholder_text="Pega aquí la clave (empieza con AIza...)",
            show="•", height=38, fg_color=COLOR_FONDO, border_color=COLOR_BORDE,
        )
        self.entrada_clave.grid(row=1, column=0, sticky="ew")
        if clave_actual:
            self.entrada_clave.insert(0, clave_actual)
        self.mostrar_clave = ctk.BooleanVar(value=False)
        ctk.CTkCheckBox(
            self, text="Mostrar", variable=self.mostrar_clave, width=90,
            command=self._alternar_clave, checkbox_width=18, checkbox_height=18,
            font=fuente_ayuda, text_color=COLOR_TEXTO_SUAVE,
        ).grid(row=1, column=1, padx=(10, 0))

        enlace = ctk.CTkLabel(
            self, text="Obtener una clave gratis en aistudio.google.com/apikey  ↗",
            font=ctk.CTkFont(size=12, underline=True), text_color=COLOR_ENLACE,
            anchor="w", cursor="hand2",
        )
        enlace.grid(row=2, column=0, columnspan=2, sticky="w", pady=(4, 0))
        enlace.bind("<Button-1>", lambda _evento: self.abrir_enlace())
        ctk.CTkLabel(
            self, text="Es gratis y no pide tarjeta. Inicia sesión con tu cuenta de Google, "
                       "pulsa \"Create API key\" y copia la clave.",
            font=fuente_ayuda, text_color=COLOR_TEXTO_SUAVE, anchor="w", justify="left",
            wraplength=520,
        ).grid(row=3, column=0, columnspan=2, sticky="w", pady=(2, 16))

        # --- Ruta del Excel -----------------------------------------------
        ctk.CTkLabel(self, text="Carpeta del Excel compartido (opcional)",
                     font=fuente_etiqueta, text_color=COLOR_TEXTO, anchor="w")\
            .grid(row=4, column=0, columnspan=2, sticky="w", pady=(0, 4))
        self.entrada_ruta = ctk.CTkEntry(
            self, placeholder_text=f"Vacío = {RUTA_EXCEL_POR_DEFECTO} (carpeta del programa)",
            height=38, fg_color=COLOR_FONDO, border_color=COLOR_BORDE,
        )
        self.entrada_ruta.grid(row=5, column=0, sticky="ew")
        if ruta_actual:
            self.entrada_ruta.insert(0, ruta_actual)
        ctk.CTkButton(
            self, text="Examinar...", command=self.examinar, width=110, height=38,
            fg_color=COLOR_TARJETA, hover_color=COLOR_TARJETA_HOVER,
            border_width=1, border_color=COLOR_BORDE, text_color=COLOR_TEXTO,
        ).grid(row=5, column=1, padx=(10, 0))
        ctk.CTkLabel(
            self, text="Elige una carpeta (por ejemplo de Google Drive) para que varios "
                       f"PCs guarden en el mismo Excel. Se usará el archivo {NOMBRE_EXCEL_POR_DEFECTO} "
                       "dentro de esa carpeta.",
            font=fuente_ayuda, text_color=COLOR_TEXTO_SUAVE, anchor="w", justify="left",
            wraplength=520,
        ).grid(row=6, column=0, columnspan=2, sticky="w", pady=(2, 18))

        # --- Guardar ------------------------------------------------------
        self.boton_guardar = ctk.CTkButton(
            self, text=texto_boton, command=self.guardar, height=44, corner_radius=22,
            font=ctk.CTkFont(size=15, weight="bold"),
            fg_color=COLOR_ACENTO, hover_color=COLOR_ACENTO_HOVER, text_color="white",
        )
        self.boton_guardar.grid(row=7, column=0, columnspan=2, sticky="ew")
        self.entrada_clave.bind("<Return>", lambda _evento: self.guardar())
        self.entrada_ruta.bind("<Return>", lambda _evento: self.guardar())

    def _alternar_clave(self):
        self.entrada_clave.configure(show="" if self.mostrar_clave.get() else "•")

    def abrir_enlace(self):
        try:
            webbrowser.open(URL_API_KEY)
        except Exception:
            registrar_error("No se pudo abrir el navegador")
            messagebox.showinfo(
                "Abrir enlace",
                f"No se pudo abrir el navegador. Entra manualmente a:\n{URL_API_KEY}",
                parent=self.winfo_toplevel(),
            )

    def examinar(self):
        try:
            actual = self.entrada_ruta.get().strip()
            inicial = os.path.dirname(actual) if actual else str(BASE)
            if not os.path.isdir(inicial):
                inicial = str(BASE)
            carpeta = filedialog.askdirectory(
                title="Elige la carpeta donde se guardará el Excel",
                initialdir=inicial, parent=self.winfo_toplevel(),
            )
        except Exception:
            registrar_error("Error en el selector de carpetas")
            carpeta = ""
        if not carpeta:
            return
        ruta = os.path.normpath(os.path.join(carpeta, NOMBRE_EXCEL_POR_DEFECTO))
        self.entrada_ruta.delete(0, "end")
        self.entrada_ruta.insert(0, ruta)

    def guardar(self):
        ventana = self.winfo_toplevel()
        clave = self.entrada_clave.get().strip()
        ruta = self.entrada_ruta.get().strip().strip('"').strip("'")

        if not api_key_valida(clave):
            messagebox.showwarning(
                "Falta la clave", "Pega la clave de API antes de guardar.", parent=ventana)
            self.entrada_clave.focus_set()
            return
        if " " in clave or "\t" in clave:
            messagebox.showwarning(
                "Clave con espacios",
                "La clave no debe tener espacios. Revisa que la hayas copiado completa.",
                parent=ventana)
            return
        if not clave.startswith("AIza"):
            if not messagebox.askyesno(
                    "Revisar la clave",
                    "Las claves de Gemini normalmente empiezan con \"AIza\".\n"
                    "¿Quieres guardarla de todos modos?", parent=ventana):
                return

        if ruta:
            # Si el usuario escribió solo una carpeta, se completa con el nombre del archivo
            if os.path.isdir(ruta) or ruta.endswith(("\\", "/")):
                ruta = os.path.join(ruta, NOMBRE_EXCEL_POR_DEFECTO)
            elif not ruta.lower().endswith(".xlsx"):
                ruta = ruta + ".xlsx"
            ruta = os.path.normpath(ruta)
        else:
            ruta = RUTA_EXCEL_POR_DEFECTO

        try:
            guardar_config(clave, ruta)
        except Exception:
            registrar_error("No se pudo guardar config.txt")
            messagebox.showerror(
                "No se pudo guardar",
                "No se pudo escribir el archivo config.txt.\n"
                "Verifica que la carpeta del programa permita guardar archivos "
                "y que config.txt no esté abierto en otro programa.",
                parent=ventana)
            return

        # Comprobar que la carpeta del Excel se pueda usar (crea las carpetas)
        try:
            leer_ruta_excel()
        except OSError:
            registrar_error("La carpeta del Excel no está disponible")
            messagebox.showwarning(
                "Carpeta del Excel",
                "La configuración se guardó, pero no se pudo acceder a la carpeta del Excel.\n"
                "Verifica que la ruta sea correcta y que la unidad (por ejemplo Google Drive) "
                "esté disponible.", parent=ventana)

        try:
            self.al_guardar()
        except Exception:
            registrar_error("Error después de guardar la configuración")


class DialogoConfiguracion(ctk.CTkToplevel):
    """Ventana ⚙️ para cambiar la clave de API y la ruta del Excel."""

    def __init__(self, padre, al_guardar):
        super().__init__(padre, fg_color=COLOR_FONDO)
        self.al_guardar = al_guardar
        self.title("Configuración")
        centrar_sobre(self, padre, 620, 470)
        self.minsize(560, 430)
        self.configure(fg_color=COLOR_FONDO)
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)

        ctk.CTkLabel(self, text="Configuración", font=ctk.CTkFont(size=22, weight="bold"),
                     text_color=COLOR_TEXTO, anchor="w")\
            .grid(row=0, column=0, sticky="w", padx=28, pady=(22, 6))
        tarjeta = ctk.CTkFrame(self, fg_color=COLOR_PANEL, corner_radius=16,
                               border_width=1, border_color=COLOR_BORDE)
        tarjeta.grid(row=1, column=0, sticky="nsew", padx=24, pady=(0, 10))
        tarjeta.grid_columnconfigure(0, weight=1)
        self.formulario = FormularioConfiguracion(tarjeta, self._guardado,
                                                  texto_boton="Guardar cambios")
        self.formulario.grid(row=0, column=0, sticky="nsew", padx=22, pady=20)
        ctk.CTkButton(
            self, text="Cancelar", command=self.destroy, width=120, height=36,
            fg_color=COLOR_TARJETA, hover_color=COLOR_TARJETA_HOVER,
            border_width=1, border_color=COLOR_BORDE, text_color=COLOR_TEXTO,
        ).grid(row=2, column=0, sticky="e", padx=24, pady=(0, 18))

        self.transient(padre)
        self.bind("<Escape>", lambda _evento: self.destroy())
        # CTkToplevel tarda unos ms en dibujarse; grab_set después de eso
        self.after(200, self._tomar_foco)

    def _tomar_foco(self):
        try:
            self.grab_set()
            self.focus_force()
            self.formulario.entrada_clave.focus_set()
        except Exception:
            pass

    def _guardado(self):
        try:
            self.al_guardar()
        finally:
            self.destroy()


# ---------------------------------------------------------------------------
# Ventana principal
# ---------------------------------------------------------------------------
class Aplicacion(ctk.CTk):
    CAMPOS = [
        ("banco_app", "BANCO / APP"),
        ("numero_comprobante", "COMPROBANTE"),
        ("numero_cuenta_o_llave", "CUENTA / LLAVE"),
        ("tipo_cuenta_o_llave", "TIPO CUENTA / LLAVE"),
        ("nombre_cliente", "NOMBRE"),
        ("valor_pago", "VALOR"),
        ("fecha_pago", "FECHA DEL PAGO"),
    ]

    # Columnas visibles del historial: (clave, título, ancho)
    COLUMNAS_HISTORIAL = [
        ("fecha_procesado", "Procesado", 110),
        ("archivo_origen", "Archivo", 140),
        ("banco_app", "Banco / App", 90),
        ("numero_comprobante", "Comprobante", 100),
        ("nombre_cliente", "Nombre", 140),
        ("valor_pago", "Valor", 90),
        ("fecha_pago", "Fecha pago", 90),
        ("estado", "Estado", 70),
    ]

    def __init__(self):
        super().__init__()
        self.procesando = False
        self.filas = {}             # nombre de archivo -> widgets de la fila
        self.seleccionados = set()  # nombres seleccionados en la lista
        self.resumen = {"total": 0, "ok": 0, "revision": 0, "errores": 0}
        self.historial_registros = {}   # id de fila del Treeview -> datos del Excel
        self.cargando_historial = False
        self.historial_pendiente = False
        self._aviso_historial_mostrado = None
        self.camara_disponible = None   # None = aún no se ha detectado
        self.ventana_camara = None
        self.dialogo_config = None
        self.marco_bienvenida = None

        self.title("Extractor de Comprobantes")
        self.geometry("1160x860")
        self.minsize(1040, 720)
        self.configure(fg_color=COLOR_FONDO)

        self.fuente_titulo = ctk.CTkFont(size=26, weight="bold")
        self.fuente_seccion = ctk.CTkFont(size=16, weight="bold")
        self.fuente_normal = ctk.CTkFont(size=13)
        self.fuente_pequena = ctk.CTkFont(size=11)
        self.fuente_etiqueta = ctk.CTkFont(size=10, weight="bold")
        self.fuente_valor = ctk.CTkFont(size=18, weight="bold")
        self.fuente_boton = ctk.CTkFont(size=13, weight="bold")
        self.fuente_procesar = ctk.CTkFont(size=18, weight="bold")
        self.fuente_contador = ctk.CTkFont(size=22, weight="bold")

        self.grid_columnconfigure(0, weight=1, uniform="paneles")
        self.grid_columnconfigure(1, weight=1, uniform="paneles")
        self.grid_rowconfigure(1, weight=1)

        self._construir_encabezado()
        self._construir_panel_izquierdo()
        self._construir_panel_derecho()
        self._construir_progreso()
        self._construir_acciones()

        self.refrescar_lista()
        self.actualizar_estado_api()
        self._actualizar_ruta_excel()
        self.refrescar_historial()
        self.after(INTERVALO_REFRESCO_MS, self._refresco_periodico)
        # La detección de cámara es lenta: se hace en un hilo tras abrir la ventana
        self.after(300, self._iniciar_deteccion_camara)

        if not leer_api_key_gui():
            self.mostrar_bienvenida()

    # Errores de Tk (callbacks de botones, after, etc.) -> log + aviso amigable
    def report_callback_exception(self, tipo, valor, tb):
        try:
            LOG.error("Error en la interfaz", exc_info=(tipo, valor, tb))
        except Exception:
            pass
        _mostrar_error_inesperado()

    def _en_ui(self, funcion, *args):
        """Ejecuta 'funcion' en el hilo principal (seguro desde otros hilos)."""
        try:
            self.after(0, funcion, *args)
        except Exception:
            pass

    # ------------------------------------------------------------------
    # Construcción de la interfaz
    # ------------------------------------------------------------------
    def _construir_encabezado(self):
        encabezado = ctk.CTkFrame(self, fg_color="transparent")
        encabezado.grid(row=0, column=0, columnspan=2, sticky="ew",
                        padx=24, pady=(18, 10))
        encabezado.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            encabezado, text="Extractor de Comprobantes",
            font=self.fuente_titulo, text_color=COLOR_TEXTO, anchor="w",
        ).grid(row=0, column=0, sticky="w")
        self.etiqueta_subtitulo = ctk.CTkLabel(
            encabezado, text=f"Modelo: {MODELO}",
            font=self.fuente_pequena, text_color=COLOR_TEXTO_SUAVE, anchor="w",
        )
        self.etiqueta_subtitulo.grid(row=1, column=0, sticky="w")

        pastilla = ctk.CTkFrame(encabezado, fg_color=COLOR_PANEL, corner_radius=18,
                                border_width=1, border_color=COLOR_BORDE)
        pastilla.grid(row=0, column=1, rowspan=2, sticky="e")
        self.punto_api = ctk.CTkLabel(pastilla, text="●", font=ctk.CTkFont(size=16),
                                      text_color=COLOR_TEXTO_SUAVE)
        self.punto_api.grid(row=0, column=0, padx=(14, 6), pady=6)
        self.texto_api = ctk.CTkLabel(pastilla, text="Verificando API key...",
                                      font=self.fuente_boton, text_color=COLOR_TEXTO)
        self.texto_api.grid(row=0, column=1, padx=(0, 16), pady=6)

        self.boton_config = self._boton_secundario(
            encabezado, "⚙️", "Configuración", self.abrir_configuracion, width=150)
        self.boton_config.grid(row=0, column=2, rowspan=2, sticky="e", padx=(12, 0))
        Tooltip(self.boton_config, "Cambiar la clave de API o la carpeta del Excel")

    def _boton_secundario(self, padre, emoji, texto, comando, **opciones):
        icono = icono_emoji(emoji)
        valores = dict(
            text=texto, command=comando, font=self.fuente_normal,
            fg_color=COLOR_TARJETA, hover_color=COLOR_TARJETA_HOVER,
            text_color=COLOR_TEXTO, corner_radius=10, height=38,
            border_width=1, border_color=COLOR_BORDE,
        )
        if icono is not None:
            valores.update(image=icono, compound="left")
        else:
            valores["text"] = f"{emoji} {texto}"
        valores.update(opciones)
        return ctk.CTkButton(padre, **valores)

    def _construir_panel_izquierdo(self):
        panel = ctk.CTkFrame(self, fg_color=COLOR_PANEL, corner_radius=16,
                             border_width=1, border_color=COLOR_BORDE)
        panel.grid(row=1, column=0, sticky="nsew", padx=(24, 8), pady=8)
        panel.grid_columnconfigure(0, weight=1)
        panel.grid_rowconfigure(1, weight=1)

        cabecera = ctk.CTkFrame(panel, fg_color="transparent")
        cabecera.grid(row=0, column=0, sticky="ew", padx=16, pady=(14, 8))
        cabecera.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(cabecera, text="Imágenes pendientes", font=self.fuente_seccion,
                     text_color=COLOR_TEXTO, anchor="w").grid(row=0, column=0, sticky="w")
        self.etiqueta_conteo = ctk.CTkLabel(
            cabecera, text="0", font=self.fuente_boton, text_color=COLOR_TEXTO,
            fg_color=COLOR_ACENTO, corner_radius=10, width=34, height=24,
        )
        self.etiqueta_conteo.grid(row=0, column=1, sticky="e")
        self.etiqueta_ayuda_lista = ctk.CTkLabel(
            cabecera, text="Haz clic en las filas para seleccionarlas",
            font=self.fuente_pequena, text_color=COLOR_TEXTO_SUAVE, anchor="w",
        )
        self.etiqueta_ayuda_lista.grid(row=1, column=0, columnspan=2, sticky="w")

        self.lista = ctk.CTkScrollableFrame(
            panel, fg_color=COLOR_FONDO, corner_radius=12,
            scrollbar_button_color=COLOR_BORDE,
            scrollbar_button_hover_color=COLOR_TARJETA_HOVER,
        )
        self.lista.grid(row=1, column=0, sticky="nsew", padx=16)
        self.lista.grid_columnconfigure(0, weight=1)

        self.etiqueta_vacia = ctk.CTkLabel(
            self.lista,
            text='No hay imágenes pendientes.\n\nUsa "Agregar imágenes", "Agregar carpeta"\n'
                 'o "Usar cámara", o guarda los comprobantes\nen la carpeta "entrada".',
            font=self.fuente_normal, text_color=COLOR_TEXTO_SUAVE, justify="center",
        )

        botones = ctk.CTkFrame(panel, fg_color="transparent")
        botones.grid(row=2, column=0, sticky="ew", padx=16, pady=14)
        for columna in range(3):
            botones.grid_columnconfigure(columna, weight=1, uniform="botones")
        self._boton_secundario(botones, "➕", "Agregar imágenes", self.agregar_imagenes)\
            .grid(row=0, column=0, sticky="ew", padx=(0, 4), pady=(0, 6))
        self._boton_secundario(botones, "🗂", "Agregar carpeta", self.agregar_carpeta)\
            .grid(row=0, column=1, sticky="ew", padx=4, pady=(0, 6))
        self.boton_camara = self._boton_secundario(
            botones, "📷", "Usar cámara", self.abrir_camara, state="disabled")
        self.boton_camara.grid(row=0, column=2, sticky="ew", padx=(4, 0), pady=(0, 6))
        self.tooltip_camara = Tooltip(self.boton_camara, "Buscando cámara web...")
        self._boton_secundario(botones, "📂", "Abrir carpeta", self.abrir_entrada)\
            .grid(row=1, column=0, sticky="ew", padx=(0, 4))
        self._boton_secundario(botones, "🗑", "Quitar selección", self.quitar_seleccion,
                               hover_color="#4a2530")\
            .grid(row=1, column=1, sticky="ew", padx=4)

    def _construir_panel_derecho(self):
        panel = ctk.CTkFrame(self, fg_color=COLOR_PANEL, corner_radius=16,
                             border_width=1, border_color=COLOR_BORDE)
        panel.grid(row=1, column=1, sticky="nsew", padx=(8, 24), pady=8)
        panel.grid_columnconfigure(0, weight=1)
        panel.grid_rowconfigure(3, weight=1)

        cabecera = ctk.CTkFrame(panel, fg_color="transparent")
        cabecera.grid(row=0, column=0, sticky="ew", padx=16, pady=(14, 8))
        cabecera.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(cabecera, text="Último comprobante", font=self.fuente_seccion,
                     text_color=COLOR_TEXTO, anchor="w").grid(row=0, column=0, sticky="w")
        self.etiqueta_archivo = ctk.CTkLabel(
            cabecera, text="Aún no se ha procesado ningún comprobante",
            font=self.fuente_pequena, text_color=COLOR_TEXTO_SUAVE, anchor="w",
        )
        self.etiqueta_archivo.grid(row=1, column=0, sticky="w")
        self.badge_estado = ctk.CTkLabel(
            cabecera, text="SIN DATOS", font=self.fuente_boton,
            text_color=COLOR_TEXTO_SUAVE, fg_color=COLOR_TARJETA,
            corner_radius=14, width=110, height=32,
        )
        self.badge_estado.grid(row=0, column=1, rowspan=2, sticky="e")

        tarjetas = ctk.CTkFrame(panel, fg_color="transparent")
        tarjetas.grid(row=1, column=0, sticky="ew", padx=10)
        tarjetas.grid_columnconfigure((0, 1), weight=1, uniform="tarjetas")

        self.campos = {}
        for indice, (clave, titulo) in enumerate(self.CAMPOS):
            tarjeta = ctk.CTkFrame(tarjetas, fg_color=COLOR_TARJETA, corner_radius=12)
            # Si la cantidad de campos es impar, la última tarjeta ocupa toda la fila.
            # Los márgenes van ajustados: son 4 filas de tarjetas y el historial
            # de abajo necesita su espacio.
            ancho = 2 if indice == len(self.CAMPOS) - 1 and indice % 2 == 0 else 1
            tarjeta.grid(row=indice // 2, column=indice % 2, columnspan=ancho,
                         sticky="nsew", padx=6, pady=4)
            tarjeta.grid_columnconfigure(0, weight=1)
            ctk.CTkLabel(tarjeta, text=titulo, font=self.fuente_etiqueta,
                         text_color=COLOR_TEXTO_SUAVE, anchor="w", height=16)\
                .grid(row=0, column=0, sticky="w", padx=14, pady=(7, 0))
            valor = ctk.CTkLabel(tarjeta, text="—", font=self.fuente_valor,
                                 text_color=COLOR_TEXTO, anchor="w", justify="left")
            valor.grid(row=1, column=0, sticky="w", padx=14, pady=(0, 6))
            tarjeta.bind("<Configure>",
                         lambda evento, etiqueta=valor: self._ajustar_ajuste(evento, etiqueta))
            self.campos[clave] = valor

        self.marco_motivos = ctk.CTkFrame(panel, fg_color="#2b2614", corner_radius=12,
                                          border_width=1, border_color="#5c4f16")
        self.marco_motivos.grid(row=2, column=0, sticky="ew", padx=16, pady=(6, 6))
        self.marco_motivos.grid_columnconfigure(0, weight=1)
        self.etiqueta_motivos = ctk.CTkLabel(
            self.marco_motivos, text="", font=self.fuente_normal,
            text_color=COLOR_REVISAR, anchor="w", justify="left",
        )
        self.etiqueta_motivos.grid(row=0, column=0, sticky="ew", padx=14, pady=10)
        self.marco_motivos.bind(
            "<Configure>",
            lambda evento: self._ajustar_ajuste(evento, self.etiqueta_motivos))
        self.marco_motivos.grid_remove()

        self._construir_historial(panel)

    def _construir_historial(self, panel):
        """Tabla con los últimos comprobantes leídos del Excel."""
        marco = ctk.CTkFrame(panel, fg_color="transparent")
        marco.grid(row=3, column=0, sticky="nsew", padx=16, pady=(6, 14))
        marco.grid_columnconfigure(0, weight=1)
        marco.grid_rowconfigure(2, weight=1)

        # El título va corto para que el contador quepa siempre a su derecha
        ctk.CTkLabel(marco, text="Historial", font=self.fuente_seccion,
                     text_color=COLOR_TEXTO, anchor="w")\
            .grid(row=0, column=0, sticky="w")
        self.etiqueta_totales = ctk.CTkLabel(
            marco, text="Total procesados hoy: 0   |   Total general: 0",
            font=self.fuente_boton, text_color=COLOR_TEXTO, anchor="e",
        )
        self.etiqueta_totales.grid(row=0, column=1, sticky="e", padx=(12, 0))
        self.AYUDA_HISTORIAL = (f"Últimos {LIMITE_HISTORIAL} comprobantes del Excel. "
                                "Doble clic en una fila abre la imagen original.")
        self.etiqueta_ayuda_historial = ctk.CTkLabel(
            marco, text=self.AYUDA_HISTORIAL,
            font=self.fuente_pequena, text_color=COLOR_TEXTO_SUAVE, anchor="w",
        )
        self.etiqueta_ayuda_historial.grid(row=1, column=0, columnspan=2, sticky="w",
                                           pady=(0, 6))

        contenedor = ctk.CTkFrame(marco, fg_color=COLOR_FONDO, corner_radius=12)
        contenedor.grid(row=2, column=0, columnspan=2, sticky="nsew")
        contenedor.grid_columnconfigure(0, weight=1)
        contenedor.grid_rowconfigure(0, weight=1)

        # Estilo oscuro para el Treeview (ttk)
        familia = self.fuente_normal.cget("family")
        estilo = ttk.Style(self)
        try:
            estilo.theme_use("clam")
        except tk.TclError:
            pass
        # El tema "clam" dibuja bordes claros (lightcolor/darkcolor); se oscurecen
        estilo.configure(
            "Historial.Treeview", background=COLOR_FONDO, fieldbackground=COLOR_FONDO,
            foreground=COLOR_TEXTO, rowheight=26, borderwidth=0, relief="flat",
            font=(familia, 10), bordercolor=COLOR_FONDO,
            lightcolor=COLOR_FONDO, darkcolor=COLOR_FONDO,
        )
        estilo.configure(
            "Historial.Treeview.Heading", background=COLOR_TARJETA,
            foreground=COLOR_TEXTO_SUAVE, relief="flat", borderwidth=0,
            font=(familia, 9, "bold"), padding=(6, 4), bordercolor=COLOR_FONDO,
            lightcolor=COLOR_TARJETA, darkcolor=COLOR_TARJETA,
        )
        estilo.map("Historial.Treeview",
                   background=[("selected", COLOR_SELECCION)],
                   foreground=[("selected", COLOR_TEXTO)])
        estilo.map("Historial.Treeview.Heading",
                   background=[("active", COLOR_TARJETA_HOVER)])
        for orientacion in ("Vertical", "Horizontal"):
            estilo.configure(
                f"Historial.{orientacion}.TScrollbar", background=COLOR_TARJETA,
                troughcolor=COLOR_FONDO, bordercolor=COLOR_FONDO,
                lightcolor=COLOR_TARJETA, darkcolor=COLOR_TARJETA,
                arrowcolor=COLOR_TEXTO_SUAVE, borderwidth=0, arrowsize=12,
            )
            estilo.map(f"Historial.{orientacion}.TScrollbar",
                       background=[("active", COLOR_TARJETA_HOVER),
                                   ("pressed", COLOR_TARJETA_HOVER)])

        claves = [clave for clave, _titulo, _ancho in self.COLUMNAS_HISTORIAL]
        self.tabla = ttk.Treeview(contenedor, columns=claves, show="headings",
                                  style="Historial.Treeview", selectmode="browse")
        for clave, titulo, ancho in self.COLUMNAS_HISTORIAL:
            self.tabla.heading(clave, text=titulo, anchor="w")
            self.tabla.column(clave, width=ancho, minwidth=50, stretch=False,
                              anchor="e" if clave == "valor_pago" else "w")
        self.tabla.tag_configure("ok", foreground=COLOR_OK)
        self.tabla.tag_configure("revisar", foreground=COLOR_REVISAR)
        self.tabla.grid(row=0, column=0, sticky="nsew", padx=(6, 0), pady=6)

        barra_v = ttk.Scrollbar(contenedor, orient="vertical", command=self.tabla.yview,
                                style="Historial.Vertical.TScrollbar")
        barra_v.grid(row=0, column=1, sticky="ns", pady=6)
        barra_h = ttk.Scrollbar(contenedor, orient="horizontal", command=self.tabla.xview,
                                style="Historial.Horizontal.TScrollbar")
        barra_h.grid(row=1, column=0, sticky="ew", padx=(6, 0))
        self.tabla.configure(yscrollcommand=barra_v.set, xscrollcommand=barra_h.set)
        self.tabla.bind("<Double-1>", self._abrir_imagen_historial)
        self.tabla.bind("<Return>", self._abrir_imagen_historial)

    def _ajustar_ajuste(self, evento, etiqueta):
        """Ajusta el wraplength de una etiqueta al ancho de su tarjeta."""
        escala = ctk.ScalingTracker.get_widget_scaling(self)
        ancho = max(80, int(evento.width / escala) - 32)
        etiqueta.configure(wraplength=ancho)

    def _construir_progreso(self):
        marco = ctk.CTkFrame(self, fg_color="transparent")
        marco.grid(row=2, column=0, columnspan=2, sticky="ew", padx=24, pady=(8, 0))
        marco.grid_columnconfigure(0, weight=1)

        self.barra = ctk.CTkProgressBar(marco, height=12, corner_radius=6,
                                        fg_color=COLOR_TARJETA, progress_color=COLOR_ACENTO)
        self.barra.grid(row=0, column=0, sticky="ew")
        self.barra.set(0)
        self.etiqueta_porcentaje = ctk.CTkLabel(marco, text="0 %", width=56, anchor="e",
                                                font=self.fuente_boton,
                                                text_color=COLOR_TEXTO)
        self.etiqueta_porcentaje.grid(row=0, column=1, padx=(12, 0))
        self.etiqueta_estado = ctk.CTkLabel(marco, text="Listo.", anchor="w",
                                            font=self.fuente_normal,
                                            text_color=COLOR_TEXTO_SUAVE)
        self.etiqueta_estado.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(2, 0))

    def _construir_acciones(self):
        marco = ctk.CTkFrame(self, fg_color="transparent")
        marco.grid(row=3, column=0, columnspan=2, sticky="ew", padx=24, pady=(6, 20))
        marco.grid_columnconfigure(2, weight=1)

        self.boton_procesar = ctk.CTkButton(
            marco, text="PROCESAR", command=self.iniciar_procesamiento,
            font=self.fuente_procesar, width=220, height=58, corner_radius=29,
            fg_color=COLOR_ACENTO, hover_color=COLOR_ACENTO_HOVER,
            text_color="white", text_color_disabled="#aeb2bf",
        )
        self.boton_procesar.grid(row=0, column=0, sticky="w")

        self._boton_secundario(
            marco, "📊", "Abrir Excel", self.abrir_excel,
            width=170, height=58, corner_radius=29, font=self.fuente_boton,
        ).grid(row=0, column=1, sticky="w", padx=(12, 0))

        resumen = ctk.CTkFrame(marco, fg_color="transparent")
        resumen.grid(row=0, column=3, sticky="e")
        self.contadores = {}
        definicion = [
            ("total", "Procesados", COLOR_TEXTO),
            ("ok", "OK", COLOR_OK),
            ("revision", "Revisión", COLOR_REVISAR),
            ("errores", "Errores", COLOR_ERROR),
        ]
        for columna, (clave, titulo, color) in enumerate(definicion):
            tarjeta = ctk.CTkFrame(resumen, fg_color=COLOR_PANEL, corner_radius=12,
                                   border_width=1, border_color=COLOR_BORDE,
                                   width=100, height=58)
            tarjeta.grid(row=0, column=columna, padx=(8, 0))
            tarjeta.grid_propagate(False)
            tarjeta.grid_columnconfigure(0, weight=1)
            numero = ctk.CTkLabel(tarjeta, text="0", font=self.fuente_contador,
                                  text_color=color, height=26)
            numero.grid(row=0, column=0, pady=(6, 0))
            ctk.CTkLabel(tarjeta, text=titulo, font=self.fuente_pequena,
                         text_color=COLOR_TEXTO_SUAVE, height=16)\
                .grid(row=1, column=0, pady=(0, 6))
            self.contadores[clave] = numero
        ctk.CTkLabel(resumen, text="Acumulado de esta sesión", font=self.fuente_pequena,
                     text_color=COLOR_TEXTO_SUAVE, height=14)\
            .grid(row=1, column=0, columnspan=4, sticky="e", pady=(4, 0))

    # ------------------------------------------------------------------
    # Pantalla de bienvenida (primera configuración)
    # ------------------------------------------------------------------
    def mostrar_bienvenida(self):
        """Cubre la ventana con la pantalla de configuración inicial."""
        if self.marco_bienvenida is not None:
            return
        self.marco_bienvenida = ctk.CTkFrame(self, fg_color=COLOR_FONDO, corner_radius=0)
        self.marco_bienvenida.grid(row=0, column=0, rowspan=4, columnspan=2, sticky="nsew")
        self.marco_bienvenida.grid_columnconfigure(0, weight=1)
        self.marco_bienvenida.grid_rowconfigure(0, weight=1)
        self.marco_bienvenida.grid_rowconfigure(2, weight=1)

        tarjeta = ctk.CTkFrame(self.marco_bienvenida, fg_color=COLOR_PANEL, corner_radius=20,
                               border_width=1, border_color=COLOR_BORDE, width=640)
        tarjeta.grid(row=1, column=0, padx=40)
        tarjeta.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(tarjeta, text="👋  Bienvenido al Extractor de Comprobantes",
                     font=self.fuente_titulo, text_color=COLOR_TEXTO, anchor="w")\
            .grid(row=0, column=0, sticky="w", padx=32, pady=(28, 4))
        ctk.CTkLabel(
            tarjeta,
            text="Antes de empezar necesitamos dos cosas: la clave de API de Google Gemini "
                 "(gratis) y, si quieres, la carpeta donde se guardará el Excel. "
                 "Todo se guarda automáticamente en config.txt; no hace falta editar ningún archivo.",
            font=self.fuente_normal, text_color=COLOR_TEXTO_SUAVE, anchor="w",
            justify="left", wraplength=560,
        ).grid(row=1, column=0, sticky="w", padx=32, pady=(0, 20))
        formulario = FormularioConfiguracion(tarjeta, self._bienvenida_guardada,
                                             texto_boton="Guardar y empezar")
        formulario.grid(row=2, column=0, sticky="ew", padx=32, pady=(0, 28))
        self.marco_bienvenida.lift()
        self.after(200, lambda: enfocar_si_existe(formulario.entrada_clave))

    def _bienvenida_guardada(self):
        if self.marco_bienvenida is not None:
            self.marco_bienvenida.destroy()
            self.marco_bienvenida = None
        self.recargar_configuracion()
        self.mostrar_estado("Configuración guardada. Ya puedes agregar imágenes y procesar.",
                            COLOR_OK)

    # ------------------------------------------------------------------
    # Configuración (⚙️)
    # ------------------------------------------------------------------
    def abrir_configuracion(self):
        if self.dialogo_config is not None and self.dialogo_config.winfo_exists():
            self.dialogo_config.lift()
            self.dialogo_config.focus_force()
            return
        try:
            self.dialogo_config = DialogoConfiguracion(self, self.recargar_configuracion)
        except Exception:
            registrar_error("No se pudo abrir la ventana de configuración")
            messagebox.showerror("Configuración",
                                 "No se pudo abrir la ventana de configuración.")

    def recargar_configuracion(self):
        """Vuelve a leer config.txt y actualiza todo lo que depende de él."""
        self.actualizar_estado_api()
        self._actualizar_ruta_excel()
        self._aviso_historial_mostrado = None
        self.refrescar_historial()

    # ------------------------------------------------------------------
    # Encabezado: estado de la API key y ruta del Excel
    # ------------------------------------------------------------------
    def actualizar_estado_api(self):
        if leer_api_key_gui():
            self.punto_api.configure(text_color=COLOR_OK)
            self.texto_api.configure(text="API key configurada")
        else:
            self.punto_api.configure(text_color=COLOR_ERROR)
            self.texto_api.configure(text="Falta API key")

    def _actualizar_ruta_excel(self):
        try:
            ruta = leer_ruta_excel()
            texto = f"Modelo: {MODELO}   ·   Excel: {acortar_centro(str(ruta))}"
        except OSError:
            texto = f"Modelo: {MODELO}   ·   Excel: carpeta no disponible (revisa la configuración ⚙️)"
        self.etiqueta_subtitulo.configure(text=texto)

    def _refresco_periodico(self):
        try:
            self.actualizar_estado_api()
            if not self.procesando:
                self.refrescar_lista()
        except Exception:
            registrar_error("Error en el refresco periódico")
        finally:
            self.after(INTERVALO_REFRESCO_MS, self._refresco_periodico)

    # ------------------------------------------------------------------
    # Lista de imágenes pendientes
    # ------------------------------------------------------------------
    def refrescar_lista(self):
        try:
            crear_carpetas()
            imagenes = listar_imagenes()
        except OSError:
            registrar_error("No se pudo leer la carpeta entrada")
            self.mostrar_estado("No se pudo leer la carpeta entrada. Verifica que exista "
                                "y que tengas permiso para usarla.", COLOR_ERROR)
            return
        nombres = [imagen.name for imagen in imagenes]
        presentes = set(nombres)

        # Quitar filas de archivos que ya no están
        for nombre in list(self.filas):
            if nombre not in presentes:
                self.filas.pop(nombre)["marco"].destroy()
                self.seleccionados.discard(nombre)

        # Crear las nuevas, actualizar tamaños y reordenar
        for posicion, imagen in enumerate(imagenes):
            try:
                tamano = tamano_legible(imagen.stat().st_size)
            except OSError:
                tamano = "—"
            fila = self.filas.get(imagen.name)
            if fila is None:
                fila = self._crear_fila(imagen.name, tamano)
                self.filas[imagen.name] = fila
            elif fila["tamano"].cget("text") != tamano:
                fila["tamano"].configure(text=tamano)
            fila["marco"].grid(row=posicion, column=0, sticky="ew", padx=4, pady=3)

        total = len(nombres)
        if total == 0:
            self.etiqueta_vacia.grid(row=0, column=0, pady=40)
        else:
            self.etiqueta_vacia.grid_remove()
        self.etiqueta_conteo.configure(text=str(total))
        self._actualizar_ayuda_lista()

    def _crear_fila(self, nombre, tamano):
        marco = ctk.CTkFrame(self.lista, fg_color=COLOR_TARJETA, corner_radius=10,
                             border_width=1, border_color=COLOR_TARJETA, height=40)
        marco.grid_columnconfigure(1, weight=1)
        variable = ctk.BooleanVar(value=False)
        casilla = ctk.CTkCheckBox(
            marco, text="", variable=variable, width=22,
            checkbox_width=18, checkbox_height=18, corner_radius=5,
            border_color=COLOR_TEXTO_SUAVE, fg_color=COLOR_SELECCION_BORDE,
            command=lambda: self._sincronizar_fila(nombre),
        )
        casilla.grid(row=0, column=0, padx=(10, 4), pady=8)
        etiqueta = ctk.CTkLabel(marco, text=acortar(nombre), font=self.fuente_normal,
                                text_color=COLOR_TEXTO, anchor="w")
        etiqueta.grid(row=0, column=1, sticky="ew")
        etiqueta_tamano = ctk.CTkLabel(marco, text=tamano, font=self.fuente_pequena,
                                       text_color=COLOR_TEXTO_SUAVE, anchor="e")
        etiqueta_tamano.grid(row=0, column=2, padx=(8, 12))

        for widget in (marco, etiqueta, etiqueta_tamano):
            widget.bind("<Button-1>", lambda _evento: self._alternar_fila(nombre))
            widget.configure(cursor="hand2")
        return {"marco": marco, "variable": variable, "tamano": etiqueta_tamano}

    def _alternar_fila(self, nombre):
        fila = self.filas.get(nombre)
        if fila is None:
            return
        fila["variable"].set(not fila["variable"].get())
        self._sincronizar_fila(nombre)

    def _sincronizar_fila(self, nombre):
        fila = self.filas.get(nombre)
        if fila is None:
            return
        if fila["variable"].get():
            self.seleccionados.add(nombre)
            fila["marco"].configure(fg_color=COLOR_SELECCION,
                                    border_color=COLOR_SELECCION_BORDE)
        else:
            self.seleccionados.discard(nombre)
            fila["marco"].configure(fg_color=COLOR_TARJETA, border_color=COLOR_TARJETA)
        self._actualizar_ayuda_lista()

    def _actualizar_ayuda_lista(self):
        if self.seleccionados:
            cantidad = len(self.seleccionados)
            texto = "1 seleccionada" if cantidad == 1 else f"{cantidad} seleccionadas"
            self.etiqueta_ayuda_lista.configure(text=texto, text_color=COLOR_SELECCION_BORDE)
        else:
            self.etiqueta_ayuda_lista.configure(
                text="Haz clic en las filas para seleccionarlas",
                text_color=COLOR_TEXTO_SUAVE)

    def agregar_imagenes(self):
        try:
            rutas = filedialog.askopenfilenames(
                title="Selecciona los comprobantes",
                filetypes=[
                    ("Imágenes", "*.jpg *.jpeg *.png *.webp"),
                    ("Todos los archivos", "*.*"),
                ],
            )
        except Exception:
            registrar_error("Error en el selector de archivos")
            rutas = ()
        if not rutas:
            return
        self._copiar_a_entrada(list(rutas))

    def agregar_carpeta(self):
        try:
            carpeta = filedialog.askdirectory(
                title="Selecciona la carpeta con los comprobantes (incluye subcarpetas)")
        except Exception:
            registrar_error("Error en el selector de carpetas")
            carpeta = ""
        if not carpeta:
            return
        try:
            rutas = buscar_imagenes_en_carpeta(carpeta)
        except OSError:
            registrar_error("No se pudo recorrer la carpeta elegida")
            messagebox.showerror("Agregar carpeta",
                                 "No se pudo leer la carpeta elegida. Verifica que exista "
                                 "y que tengas permiso para abrirla.")
            return
        if not rutas:
            messagebox.showinfo(
                "Agregar carpeta",
                "En esa carpeta (y sus subcarpetas) no hay imágenes jpg, png o webp.")
            return
        if len(rutas) > 50 and not messagebox.askyesno(
                "Agregar carpeta",
                f"Se encontraron {len(rutas)} imágenes. ¿Agregarlas todas a la lista "
                "de pendientes?"):
            return
        self._copiar_a_entrada(rutas)

    def _copiar_a_entrada(self, rutas):
        """Copia las imágenes a entrada/ sin sobreescribir las que ya existen."""
        try:
            crear_carpetas()
        except OSError:
            registrar_error("No se pudo crear la carpeta entrada")
            messagebox.showerror("Error", "No se pudo crear la carpeta entrada.")
            return
        copiadas = 0
        omitidas = 0
        fallidas = []
        for ruta in rutas:
            if os.path.splitext(ruta)[1].lower() not in EXTENSIONES:
                omitidas += 1
                continue
            try:
                # Si el archivo ya está en entrada/, no hay nada que copiar
                if os.path.exists(ruta) and os.path.normcase(
                        os.path.abspath(os.path.dirname(ruta))) == os.path.normcase(
                        os.path.abspath(str(CARPETA_ENTRADA))):
                    continue
                destino = nombre_libre(CARPETA_ENTRADA, os.path.basename(ruta))
                shutil.copy2(ruta, destino)
                copiadas += 1
            except OSError:
                registrar_error(f"No se pudo copiar {ruta}")
                fallidas.append(os.path.basename(ruta))
        self.refrescar_lista()
        if fallidas:
            detalle = "\n".join(fallidas[:8])
            if len(fallidas) > 8:
                detalle += f"\n... y {len(fallidas) - 8} más"
            messagebox.showerror(
                "No se pudieron copiar algunas imágenes",
                "Estas imágenes no se pudieron copiar a la carpeta entrada "
                f"(¿archivo abierto o sin permiso?):\n\n{detalle}")
        partes = []
        if copiadas:
            partes.append(f"Se agregaron {copiadas} imágenes a la carpeta entrada.")
        if omitidas:
            partes.append(f"{omitidas} archivos no eran imágenes jpg, png o webp.")
        if partes:
            self.mostrar_estado("  ".join(partes))

    def abrir_entrada(self):
        try:
            crear_carpetas()
            abrir_en_sistema(CARPETA_ENTRADA)
        except Exception:
            registrar_error("No se pudo abrir la carpeta entrada")
            messagebox.showerror("Error", "No se pudo abrir la carpeta entrada.")

    def quitar_seleccion(self):
        if self.procesando:
            messagebox.showinfo("Quitar selección",
                                "Espera a que termine el procesamiento.")
            return
        nombres = sorted(self.seleccionados)
        if not nombres:
            messagebox.showinfo(
                "Quitar selección",
                "Primero selecciona en la lista las imágenes que quieres quitar.",
            )
            return
        if not messagebox.askyesno(
            "Confirmar",
            f"¿Eliminar {len(nombres)} imagen(es) de la carpeta entrada?\n"
            "Esta acción no se puede deshacer.",
        ):
            return
        for nombre in nombres:
            try:
                (CARPETA_ENTRADA / nombre).unlink(missing_ok=True)
            except OSError:
                registrar_error(f"No se pudo eliminar {nombre}")
                messagebox.showerror(
                    "Error",
                    f"No se pudo eliminar {nombre}.\n"
                    "Verifica que no esté abierta en otro programa.")
        self.refrescar_lista()
        self.mostrar_estado(f"Se quitaron {len(nombres)} imágenes.")

    # ------------------------------------------------------------------
    # Cámara web
    # ------------------------------------------------------------------
    def _iniciar_deteccion_camara(self):
        threading.Thread(target=self._detectar_camara_en_hilo, daemon=True).start()

    def _detectar_camara_en_hilo(self):
        try:
            disponible, motivo = detectar_camara()
        except Exception:
            registrar_error("Error inesperado al detectar la cámara")
            disponible, motivo = False, "No se pudo comprobar si hay cámara."
        self._en_ui(self._aplicar_estado_camara, disponible, motivo)

    def _aplicar_estado_camara(self, disponible, motivo):
        self.camara_disponible = disponible
        if disponible:
            self.boton_camara.configure(state="normal", text="Usar cámara")
            self.tooltip_camara.configurar(
                "Toma una foto del comprobante con la cámara web.\n"
                "Se guarda en la carpeta entrada como camara_FECHA_HORA.jpg.")
        else:
            self.boton_camara.configure(state="disabled", text="Sin cámara")
            self.tooltip_camara.configurar(motivo)

    def abrir_camara(self):
        if not self.camara_disponible:
            messagebox.showinfo("Cámara web", "La cámara no está disponible en este equipo.")
            return
        if self.ventana_camara is not None and self.ventana_camara.winfo_exists():
            self.ventana_camara.lift()
            self.ventana_camara.focus_force()
            return
        try:
            import cv2
        except Exception:
            registrar_error("No se pudo cargar OpenCV al abrir la cámara")
            messagebox.showerror("Cámara web",
                                 "No se pudo cargar el componente de cámara (opencv-python).")
            return
        try:
            self.ventana_camara = VentanaCamara(self, cv2, self._foto_capturada)
        except Exception:
            registrar_error("No se pudo abrir la ventana de la cámara")
            messagebox.showerror("Cámara web", "No se pudo abrir la ventana de la cámara.")

    def _foto_capturada(self, ruta):
        self.refrescar_lista()
        self.mostrar_estado(f"Foto de la cámara guardada como {ruta.name}.", COLOR_OK)

    # ------------------------------------------------------------------
    # Historial de comprobantes (leído del Excel)
    # ------------------------------------------------------------------
    def refrescar_historial(self):
        """Lee el Excel en un hilo y actualiza la tabla y los totales."""
        if self.cargando_historial:
            self.historial_pendiente = True
            return
        try:
            ruta = leer_ruta_excel()
        except OSError:
            registrar_error("Ruta del Excel no disponible al leer el historial")
            self._mostrar_historial([], 0, 0,
                                    "No se pudo acceder a la carpeta del Excel.")
            return
        self.cargando_historial = True
        threading.Thread(target=self._cargar_historial_en_hilo, args=(ruta,),
                         daemon=True).start()

    def _cargar_historial_en_hilo(self, ruta):
        filas, hoy, total, aviso = [], 0, 0, None
        try:
            filas, hoy, total = leer_historial_excel(ruta)
        except PermissionError:
            registrar_error("Excel bloqueado al leer el historial")
            aviso = MENSAJE_EXCEL_NO_LEIDO
        except Exception:
            registrar_error("No se pudo leer el historial del Excel")
            aviso = MENSAJE_EXCEL_NO_LEIDO
        self._en_ui(self._mostrar_historial, filas, hoy, total, aviso)

    def _mostrar_historial(self, filas, hoy, total, aviso=None):
        self.cargando_historial = False
        try:
            self.tabla.delete(*self.tabla.get_children())
            self.historial_registros.clear()
            for registro in filas:
                valores = []
                for clave, _titulo, _ancho in self.COLUMNAS_HISTORIAL:
                    valor = registro.get(clave)
                    if clave == "valor_pago":
                        texto = formatear_moneda(valor) if valor not in (None, "") else ""
                    elif clave == "fecha_procesado":
                        texto = _texto_fecha(valor)
                    else:
                        texto = "" if valor is None else str(valor)
                    valores.append(texto)
                estado = str(registro.get("estado") or "").upper()
                etiqueta = "ok" if estado == "OK" else "revisar" if estado else ""
                identificador = self.tabla.insert("", "end", values=valores,
                                                  tags=(etiqueta,) if etiqueta else ())
                self.historial_registros[identificador] = registro
            self.etiqueta_totales.configure(
                text=f"Total procesados hoy: {hoy}   |   Total general: {total}")
            if aviso:
                self.etiqueta_ayuda_historial.configure(text=aviso, text_color=COLOR_REVISAR)
                if self._aviso_historial_mostrado != aviso:
                    self._aviso_historial_mostrado = aviso
                    messagebox.showerror("Historial", aviso)
            else:
                self.etiqueta_ayuda_historial.configure(
                    text=self.AYUDA_HISTORIAL, text_color=COLOR_TEXTO_SUAVE)
        except Exception:
            registrar_error("No se pudo mostrar el historial")
        if self.historial_pendiente:
            self.historial_pendiente = False
            self.refrescar_historial()

    def _abrir_imagen_historial(self, _evento=None):
        seleccion = self.tabla.selection()
        if not seleccion:
            return
        registro = self.historial_registros.get(seleccion[0], {})
        nombre = registro.get("archivo_origen")
        if not nombre:
            messagebox.showinfo("Abrir imagen",
                                "Esta fila del Excel no tiene nombre de archivo.")
            return
        try:
            ruta = localizar_imagen_procesada(nombre)
        except Exception:
            registrar_error("Error al buscar la imagen del historial")
            ruta = None
        if ruta is None:
            messagebox.showinfo(
                "Imagen no encontrada",
                f"La imagen \"{nombre}\" ya no está en la carpeta procesados "
                "(ni en revision_manual).\nPuede que se haya movido o borrado.")
            return
        try:
            abrir_en_sistema(ruta)
        except Exception:
            registrar_error(f"No se pudo abrir {ruta}")
            messagebox.showerror("Abrir imagen",
                                 f"No se pudo abrir la imagen {ruta.name} con el visor "
                                 "del sistema.")

    # ------------------------------------------------------------------
    # Panel derecho y pie
    # ------------------------------------------------------------------
    def mostrar_resultado(self, fila, motivos):
        self.etiqueta_archivo.configure(text=f"Archivo: {fila['archivo_origen'] or '—'}")
        for clave, _titulo in self.CAMPOS:
            if clave == "valor_pago":
                texto = formatear_moneda(fila["valor_pago"])
            else:
                texto = fila[clave] or "—"
            self.campos[clave].configure(text=texto)

        if fila["estado"] == "OK":
            self.badge_estado.configure(text="✔  OK", fg_color=COLOR_OK,
                                        text_color="#06240f")
        else:
            self.badge_estado.configure(text="⚠  REVISAR", fg_color=COLOR_REVISAR,
                                        text_color="#2b2200")

        if motivos:
            self.etiqueta_motivos.configure(
                text="Motivos:\n" + "\n".join(f"•  {motivo}" for motivo in motivos))
            self.marco_motivos.grid()
        else:
            self.etiqueta_motivos.configure(text="")
            self.marco_motivos.grid_remove()

    def mostrar_estado(self, texto, color=COLOR_TEXTO_SUAVE):
        self.etiqueta_estado.configure(text=texto, text_color=color)

    def actualizar_progreso(self, hechos, total):
        fraccion = hechos / total if total else 0
        self.barra.set(fraccion)
        self.etiqueta_porcentaje.configure(text=f"{round(fraccion * 100)} %")

    def sumar_resumen(self, clave):
        self.resumen[clave] += 1
        self.resumen["total"] += 1
        for nombre, etiqueta in self.contadores.items():
            etiqueta.configure(text=str(self.resumen[nombre]))

    def _poner_boton_procesando(self, activo):
        if activo:
            self.boton_procesar.configure(state="disabled", text="Procesando...",
                                          fg_color=COLOR_DESHABILITADO)
        else:
            self.boton_procesar.configure(state="normal", text="PROCESAR",
                                          fg_color=COLOR_ACENTO)

    # ------------------------------------------------------------------
    # Procesamiento
    # ------------------------------------------------------------------
    def iniciar_procesamiento(self):
        if self.procesando:
            return

        self.actualizar_estado_api()

        if genai is None:
            messagebox.showerror(
                "Faltan dependencias",
                "No está instalada la librería de Google Gemini.\n"
                "Solución: haz doble clic en INSTALAR.bat (o INSTALAR.command en Mac) "
                "y espera a que termine.",
            )
            return

        clave = leer_api_key_gui()
        if not clave:
            if messagebox.askyesno(
                "Falta la clave de API",
                "Todavía no hay una clave de API configurada.\n\n"
                "Se obtiene gratis (sin tarjeta) en https://aistudio.google.com/apikey\n\n"
                "¿Quieres configurarla ahora?",
            ):
                self.abrir_configuracion()
            return

        try:
            imagenes = listar_imagenes()
        except OSError:
            registrar_error("No se pudo leer la carpeta entrada")
            messagebox.showerror("Error", "No se pudo leer la carpeta entrada.")
            return
        if not imagenes:
            messagebox.showinfo(
                "Sin imágenes",
                'No hay imágenes para procesar.\n\nUsa "Agregar imágenes", "Agregar carpeta" '
                'o "Usar cámara", o guarda los comprobantes (jpg, png o webp) en la '
                'carpeta "entrada".',
            )
            return

        self.procesando = True
        self._poner_boton_procesando(True)
        self.actualizar_progreso(0, len(imagenes))
        self.mostrar_estado(f"Preparando {len(imagenes)} imágenes...")
        hilo = threading.Thread(
            target=self._procesar_en_hilo, args=(clave, imagenes), daemon=True,
        )
        hilo.start()

    def _procesar_en_hilo(self, clave, imagenes):
        """Se ejecuta en un hilo aparte. Toda actualización de la interfaz
        se envía al hilo principal con self.after."""
        total = len(imagenes)
        total_ok = total_revisar = total_error = 0
        ruta = None
        aviso = None  # (título, mensaje) de un error que detuvo el proceso

        en_ui = self._en_ui

        try:
            ruta = leer_ruta_excel()
            comprobantes_existentes = cargar_comprobantes_existentes(ruta)
            cliente = genai.Client(api_key=clave)

            for indice, archivo in enumerate(imagenes, start=1):
                en_ui(self.mostrar_estado,
                      f"[{indice}/{total}] Procesando {archivo.name} ...")
                try:
                    try:
                        fila, motivos = self._procesar_imagen(
                            cliente, archivo, ruta, comprobantes_existentes)
                    except Exception as error:
                        # Límite por minuto de Gemini: se espera y se reintenta una vez.
                        # Si vuelve a fallar, el except de abajo detiene el proceso.
                        if (isinstance(error, ExcelBloqueadoError)
                                or clasificar_error_gemini(error) != "cuota"):
                            raise
                        registrar_error(f"Límite de Gemini con {archivo.name}; se reintenta")
                        en_ui(self.mostrar_estado,
                              f"[{indice}/{total}] Se alcanzó el límite de solicitudes "
                              f"de Gemini. Esperando {ESPERA_CUOTA_SEGUNDOS} segundos "
                              f"para reintentar {archivo.name} ...", COLOR_REVISAR)
                        time.sleep(ESPERA_CUOTA_SEGUNDOS)
                        en_ui(self.mostrar_estado,
                              f"[{indice}/{total}] Reintentando {archivo.name} ...")
                        fila, motivos = self._procesar_imagen(
                            cliente, archivo, ruta, comprobantes_existentes)

                    if fila["estado"] == "OK":
                        total_ok += 1
                        destino, clave_resumen = CARPETA_PROCESADOS, "ok"
                    else:
                        total_revisar += 1
                        destino, clave_resumen = CARPETA_REVISION, "revision"

                    en_ui(self.mostrar_resultado, fila, motivos)
                    en_ui(self.sumar_resumen, clave_resumen)

                    try:
                        mover_imagen(archivo, destino)
                    except OSError:
                        registrar_error(f"No se pudo mover {archivo.name}")
                        en_ui(self.mostrar_estado,
                              f"{archivo.name} se guardó en el Excel, pero no se pudo "
                              "mover la imagen (¿está abierta en otro programa?).",
                              COLOR_REVISAR)

                except ExcelBloqueadoError as error:
                    aviso = self._aviso_excel_bloqueado(error, ruta)
                    break

                except Exception as error:
                    registrar_error(f"Error al procesar {archivo.name}")
                    tipo = clasificar_error_gemini(error)
                    if tipo in self.AVISOS_DETENCION:
                        # La imagen sigue en entrada; no tiene sentido seguir
                        aviso = self.AVISOS_DETENCION[tipo]
                        break
                    total_error += 1
                    en_ui(self.sumar_resumen, "errores")
                    if tipo == "servicio":
                        motivo = ("Gemini está saturado en este momento; "
                                  "vuelve a procesar en unos minutos")
                    elif isinstance(error, OSError):
                        motivo = "no se pudo leer o mover el archivo"
                    elif isinstance(error, ValueError):
                        motivo = ("Gemini no devolvió datos legibles para esta imagen; "
                                  "revisa que sea un comprobante nítido")
                    else:
                        motivo = ("no se pudo procesar la imagen; se reintentará "
                                  "la próxima vez")
                    en_ui(self.mostrar_estado,
                          f"Error con {archivo.name}: {motivo}.", COLOR_ERROR)

                en_ui(self.actualizar_progreso, indice, total)
                en_ui(self.refrescar_lista)

                # Pausa para respetar el límite de solicitudes por minuto
                if indice < total:
                    time.sleep(PAUSA_SEGUNDOS)

        except ExcelBloqueadoError as error:
            aviso = self._aviso_excel_bloqueado(error, ruta)
        except OSError:
            registrar_error("No se pudo acceder a la carpeta del Excel")
            aviso = (
                "No se puede acceder al Excel",
                "No se pudo acceder a la carpeta del archivo Excel.\n"
                "Revisa la ruta con el botón ⚙️ Configuración y que la unidad "
                "(por ejemplo Google Drive) esté disponible.",
            )
        except Exception as error:
            registrar_error("Error inesperado durante el procesamiento")
            aviso = self.AVISOS_DETENCION.get(clasificar_error_gemini(error)) or (
                "Error inesperado",
                "Ocurrió un error inesperado durante el procesamiento.\n"
                "Revisa la conexión a internet y vuelve a intentarlo. "
                "Las imágenes no procesadas siguen en la carpeta entrada.\n\n"
                "Si el problema se repite, envía el archivo errores.log "
                "(está en la carpeta del programa) para revisarlo.",
            )

        en_ui(self._terminar, total_ok, total_revisar, total_error, ruta, aviso)

    # Errores que detienen el procesamiento: (título, mensaje) del aviso en español
    AVISOS_DETENCION = {
        "api_key": (
            "Clave de API no válida",
            "Gemini rechazó la clave de API.\n\n"
            "Revísala con el botón ⚙️ Configuración. Si la clave fue "
            "revocada, genera una nueva gratis en\n"
            "https://aistudio.google.com/apikey",
        ),
        "conexion": (
            "Sin conexión a internet",
            "No se pudo conectar con Gemini.\n\n"
            "Revisa que el equipo tenga internet (y que ningún antivirus o "
            "proxy bloquee el programa) y vuelve a pulsar PROCESAR.\n"
            "Las imágenes que faltan siguen en la carpeta entrada.",
        ),
        "cuota": (
            "Límite de Gemini alcanzado",
            "Gemini rechazó las solicitudes por haber alcanzado el límite de uso "
            "gratuito (por minuto o por día).\n\n"
            "Espera unos minutos y vuelve a pulsar PROCESAR. Si el aviso se "
            "repite, es probable que se haya agotado el cupo del día: inténtalo "
            "mañana.\nLas imágenes que faltan siguen en la carpeta entrada.",
        ),
    }

    @staticmethod
    def _procesar_imagen(cliente, archivo, ruta, comprobantes_existentes):
        """Envía una imagen a Gemini, valida y guarda la fila en el Excel.
        Devuelve (fila, motivos). Cualquier error se propaga al llamador."""
        imagen_bytes = preparar_imagen(archivo)
        datos = extraer_datos(cliente, imagen_bytes)
        estado, motivos = validar(datos, comprobantes_existentes)
        fila = construir_fila(archivo, datos, estado)
        # Primero se guarda en Excel; si falla, la imagen queda en entrada
        agregar_y_guardar(fila, ruta)
        if fila["numero_comprobante"]:
            comprobantes_existentes.add(fila["numero_comprobante"])
        return fila, motivos

    @staticmethod
    def _aviso_excel_bloqueado(error, ruta):
        ruta_excel = getattr(error, "ruta", None) or ruta
        mensaje = "Cierra el archivo Excel y vuelve a intentar"
        if ruta_excel:
            mensaje += f"\n\nArchivo: {ruta_excel}"
        return ("Excel abierto", mensaje)

    def _terminar(self, total_ok, total_revisar, total_error, ruta, aviso=None):
        self.procesando = False
        self._poner_boton_procesando(False)
        self.refrescar_lista()
        self.actualizar_estado_api()
        self.refrescar_historial()

        resumen = (f"OK: {total_ok}   |   Para revisar: {total_revisar}"
                   f"   |   Con error: {total_error}")
        if aviso:
            titulo, mensaje = aviso
            self.mostrar_estado(f"Detenido: {titulo}.   {resumen}", COLOR_ERROR)
            messagebox.showerror(titulo, mensaje)
            return

        self.mostrar_estado(f"Terminado.   {resumen}", COLOR_TEXTO)
        destino = f"Los resultados están en:\n{ruta}" if ruta else ""
        messagebox.showinfo(
            "Proceso terminado",
            f"Procesados correctamente (OK): {total_ok}\n"
            f"Para revisión manual: {total_revisar}\n"
            f"Con error (quedan en entrada): {total_error}\n\n"
            f"{destino}",
        )

    # ------------------------------------------------------------------
    # Excel de resultados
    # ------------------------------------------------------------------
    def abrir_excel(self):
        try:
            ruta = leer_ruta_excel()
        except OSError:
            registrar_error("No se pudo acceder a la carpeta del Excel")
            messagebox.showerror(
                "No se puede acceder al Excel",
                "No se pudo acceder a la carpeta del archivo Excel.\n"
                "Revisa la ruta con el botón ⚙️ Configuración y que la unidad "
                "(por ejemplo Google Drive) esté disponible.")
            return
        if not ruta.exists():
            messagebox.showinfo(
                "Sin resultados",
                "Todavía no existe el archivo de resultados.\n"
                "Procesa al menos un comprobante primero.\n\n"
                f"Ruta configurada: {ruta}",
            )
            return
        try:
            abrir_en_sistema(ruta)
        except Exception:
            registrar_error("No se pudo abrir el Excel")
            messagebox.showerror("Error", "No se pudo abrir el archivo de Excel con el "
                                          "programa predeterminado.")


def main():
    global _app_actual
    try:
        crear_carpetas()
    except OSError:
        registrar_error("No se pudieron crear las carpetas de trabajo")
    app = Aplicacion()
    _app_actual = app
    app.mainloop()


if __name__ == "__main__":
    main()
